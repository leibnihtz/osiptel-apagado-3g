#!/usr/bin/env python3
"""
build_region_classification.py
==============================
Genera un CSV con la clasificación de región natural (Costa/Sierra/Selva)
para cada distrito del Perú, a partir de los archivos xlsx del CPV 2017 (INEI).

Fuente: INEI - Censos Nacionales de Población y Vivienda 2017
        Tabla: Centros Poblados por piso altitudinal y altitud

Regla de asignación:
  - Para cada distrito, se identifica el centro poblado (CP) con mayor
    población censada total.
  - Se toma el piso altitudinal y la altitud de ese CP como representativos
    del distrito. Justificación: el CP más poblado concentra la mayor parte
    de los usuarios móviles y, por tanto, de las mediciones de OSIPTEL.
  - El piso altitudinal (8 pisos de Pulgar Vidal) se colapsa a 3 macro-regiones:
      Costa  ← Chala, Yunga marítima
      Sierra ← Yunga fluvial, Quechua, Suni, Puna, Janca
      Selva  ← Rupa Rupa, Omagua

Uso:
  python scripts/build_region_classification.py \\
      --input-dir data/inei_cpv2017 \\
      --output data/inei_cpv2017/region_classification_by_ubigeo.csv

Requisitos: openpyxl, pandas
"""

import argparse
import glob
import os
import sys
import re

import openpyxl
import pandas as pd


# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

PISO_TO_REGION3 = {
    "Chala":            "Costa",
    "Costa o Chala":    "Costa",   # variante observada en algunos archivos INEI
    "Yunga marítima":   "Costa",
    "Yunga maritima":   "Costa",   # sin tilde, por seguridad
    "Yunga fluvial":    "Sierra",
    "Quechua":          "Sierra",
    "Suni":             "Sierra",
    "Jalca":            "Sierra",  # sinónimo de Suni en algunas fuentes
    "Puna":             "Sierra",
    "Janca":            "Sierra",
    "Cordillera":       "Sierra",  # sinónimo de Janca
    "Rupa Rupa":        "Selva",
    "Selva alta":       "Selva",   # variante
    "Omagua":           "Selva",
    "Selva baja":       "Selva",   # variante
}

# Pisos esperados (los 8 canónicos de Pulgar Vidal + variantes documentadas)
PISOS_CONOCIDOS = set(PISO_TO_REGION3.keys())


def parse_inei_number(val) -> int:
    """Convierte un valor INEI ('125 018', '-', '', None, 0) a int."""
    if val is None:
        return 0
    s = str(val).strip()
    if s in ("", "-"):
        return 0
    # quitar espacios usados como separador de miles
    s = s.replace(" ", "").replace("\u00a0", "")
    try:
        return int(float(s))
    except ValueError:
        return 0


def is_footer_row(code_val) -> bool:
    """Detecta filas de notas al pie del archivo INEI."""
    if code_val is None:
        return False
    s = str(code_val).strip()
    return s.startswith("1/") or s.startswith("2/") or s.startswith("Fuente:")


def classify_code(code_str: str, depto_prefix: str):
    """
    Clasifica una fila INEI por su código.
    Retorna: ('depto', code) | ('provincia', code) | ('distrito', code) |
             ('cp', code) | ('skip', code)
    """
    if not code_str:
        return ("skip", "")
    code = code_str.strip()
    L = len(code)

    if L == 2:
        return ("depto", code)
    elif L == 6:
        return ("distrito", code)
    elif L == 4:
        # Provincia si empieza con el prefijo del departamento
        if code.startswith(depto_prefix):
            return ("provincia", code)
        else:
            return ("cp", code)
    elif L < 4:
        # Códigos cortos (1–3 dígitos): CPs con leading zeros perdidos al leer Excel
        return ("cp", code)
    else:
        return ("skip", code)


def parse_one_file(filepath: str) -> list[dict]:
    """
    Parsea un xlsx INEI del CPV 2017 y retorna una lista de dicts, uno por
    distrito, con el CP de mayor población y su piso/altitud.
    """
    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # --- Detectar prefijo de departamento ---
    # Estrategia: buscar primero código de 2 dígitos (caso normal: '21', '07').
    # Si no existe, buscar código de 4 dígitos terminado en '00' (caso Lima: '1500').
    # Como fallback final, extraer los primeros 2 dígitos del primer código de
    # 6 dígitos (distrito) encontrado en el archivo.
    depto_prefix = None
    depto_name = None

    for r in rows:
        code = str(r[0]).strip() if r[0] else ""
        if not code or not code.isdigit() or code == "CÓDIGO":
            continue
        if len(code) == 2:
            depto_prefix = code
            depto_name = str(r[1]).strip() if r[1] else ""
            break
        if len(code) == 4 and code.endswith("00"):
            # Lima Región/Metropolitana: código '1500'
            depto_prefix = code[:2]
            depto_name = str(r[1]).strip() if r[1] else ""
            break

    # Fallback: extraer de primer código de distrito (6 dígitos)
    if depto_prefix is None:
        for r in rows:
            code = str(r[0]).strip() if r[0] else ""
            if code.isdigit() and len(code) == 6:
                depto_prefix = code[:2]
                depto_name = str(r[1]).strip() if r[1] else ""
                break

    if depto_prefix is None:
        print(f"  WARN: no se encontró código de departamento en {filepath}")
        return []

    # Limpiar nombre del departamento
    depto_name = re.sub(
        r"^(DEPARTAMENTO|PROVINCIA CONSTITUCIONAL DEL?|REGIÓN|DISTRITO)\s+",
        "", depto_name,
    )

    # --- Recorrer filas construyendo el contexto jerárquico ---
    results = []
    current_distrito_ubigeo = None
    current_distrito_name = None
    current_distrito_pob = 0
    current_provincia_name = None
    current_cps = []  # lista de (pob, piso, altitud, nombre_cp) para el distrito actual

    def flush_distrito():
        """Cierra el distrito actual: selecciona CP con mayor población."""
        nonlocal current_distrito_ubigeo, current_cps
        if current_distrito_ubigeo is None or not current_cps:
            return
        # Ordenar por población descendente; desempate por viviendas no aplica
        # porque solo guardamos población. Si empatan, el primero (orden del INEI).
        best = max(current_cps, key=lambda x: x[0])
        pob_cp, piso, altitud, nombre_cp = best

        region3 = None
        if piso and piso in PISO_TO_REGION3:
            region3 = PISO_TO_REGION3[piso]
        elif piso:
            # Piso no reconocido → fallo ruidoso
            print(
                f"  ERROR: piso altitudinal no reconocido: '{piso}' "
                f"en distrito {current_distrito_ubigeo} ({current_distrito_name}), "
                f"CP '{nombre_cp}'. Revisar PISO_TO_REGION3.",
                file=sys.stderr,
            )
            sys.exit(1)

        results.append({
            "ubigeo": current_distrito_ubigeo,
            "departamento": depto_name,
            "provincia": current_provincia_name or "",
            "distrito": current_distrito_name or "",
            "poblacion_distrito": current_distrito_pob,
            "cp_referencia": nombre_cp,
            "cp_poblacion": pob_cp,
            "altitud_m": altitud,
            "region_natural_piso": piso or "",
            "region3": region3 or "",
        })
        current_cps = []

    for r in rows:
        code_raw = str(r[0]).strip() if r[0] else ""

        # Saltar filas vacías, encabezados y footer
        if not code_raw or code_raw == "" or code_raw == "CÓDIGO":
            continue
        if is_footer_row(r[0]):
            continue

        level, code = classify_code(code_raw, depto_prefix)

        if level == "depto":
            continue

        elif level == "provincia":
            # Flush distrito previo antes de cambiar provincia
            flush_distrito()
            current_distrito_ubigeo = None
            prov_name = str(r[1]).strip() if r[1] else ""
            current_provincia_name = re.sub(r"^PROVINCIA\s+(?:DE\s+)?", "", prov_name)

        elif level == "distrito":
            # Flush distrito previo
            flush_distrito()
            current_distrito_ubigeo = code
            dist_name = str(r[1]).strip() if r[1] else ""
            current_distrito_name = re.sub(r"^DISTRITO\s+", "", dist_name)
            current_distrito_pob = parse_inei_number(r[4])
            current_cps = []

        elif level == "cp":
            if current_distrito_ubigeo is None:
                continue
            pob = parse_inei_number(r[4])
            piso = str(r[2]).strip() if r[2] else None
            if piso == "":
                piso = None
            altitud = parse_inei_number(r[3])
            nombre_cp = str(r[1]).strip() if r[1] else ""
            current_cps.append((pob, piso, altitud, nombre_cp))

    # Flush último distrito
    flush_distrito()

    return results


def main():
    parser = argparse.ArgumentParser(
        description="Clasifica distritos del Perú por región natural (CPV 2017)"
    )
    parser.add_argument(
        "--input-dir",
        default="data/inei_cpv2017",
        help="Carpeta con los archivos dptoNN.xlsx del INEI (default: data/inei_cpv2017)",
    )
    parser.add_argument(
        "--output",
        default="data/inei_cpv2017/region_classification_by_ubigeo.csv",
        help="Ruta del CSV de salida (default: data/inei_cpv2017/region_classification_by_ubigeo.csv)",
    )
    args = parser.parse_args()

    # --- Buscar archivos ---
    pattern = os.path.join(args.input_dir, "dpto*.xlsx")
    files = sorted(glob.glob(pattern))
    if not files:
        print(f"ERROR: no se encontraron archivos dpto*.xlsx en {args.input_dir}")
        sys.exit(1)
    print(f"Encontrados {len(files)} archivos xlsx en {args.input_dir}")

    # --- Parsear todos ---
    all_districts = []
    for f in files:
        fname = os.path.basename(f)
        records = parse_one_file(f)
        print(f"  {fname}: {len(records)} distritos")
        all_districts.extend(records)

    # --- Construir DataFrame ---
    df = pd.DataFrame(all_districts)

    # Forzar ubigeo como string de 6 dígitos (preservar cero inicial: Callao = 07xxxx)
    df["ubigeo"] = df["ubigeo"].astype(str).str.zfill(6)

    # --- Validaciones ---
    n_total = len(df)
    n_unique_ubigeo = df["ubigeo"].nunique()
    print(f"\nTotal distritos: {n_total}")
    print(f"Ubigeos únicos:  {n_unique_ubigeo}")

    if n_total != n_unique_ubigeo:
        dupes = df[df.duplicated(subset=["ubigeo"], keep=False)]
        print(f"  WARN: {n_total - n_unique_ubigeo} ubigeos duplicados:")
        print(dupes[["ubigeo", "departamento", "distrito"]].to_string(index=False))

    # Rango esperado: Perú tiene ~1 874 distritos (CPV 2017)
    if not (1800 <= n_total <= 1900):
        print(
            f"  WARN: se esperaban ~1874 distritos, se obtuvieron {n_total}. "
            f"Verificar cobertura de archivos."
        )

    # Conteo por región
    print("\nDistribución por región natural:")
    print(df["region3"].value_counts().to_string())

    # Distritos sin región (CP sin piso)
    sin_region = df[df["region3"] == ""]
    if len(sin_region) > 0:
        print(f"\n  WARN: {len(sin_region)} distritos sin región asignada:")
        print(
            sin_region[["ubigeo", "departamento", "distrito", "cp_referencia"]].to_string(
                index=False
            )
        )

    # Pisos encontrados
    pisos_found = df["region_natural_piso"].unique()
    print(f"\nPisos altitudinales encontrados: {sorted(p for p in pisos_found if p)}")

    unknown = set(p for p in pisos_found if p and p not in PISOS_CONOCIDOS)
    if unknown:
        print(f"  ERROR: pisos no reconocidos: {unknown}")
        sys.exit(1)

    # --- Guardar ---
    os.makedirs(os.path.dirname(args.output), exist_ok=True)
    df.to_csv(args.output, index=False, encoding="utf-8")
    print(f"\nGuardado: {args.output} ({len(df)} filas)")


if __name__ == "__main__":
    main()
